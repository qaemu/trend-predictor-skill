#!/usr/bin/env python3
"""TREND-PREDICTOR forecasting engine.

Projects constraint curves forward with calibrated uncertainty (Farmer & Lafond 2016
random walk with drift, autocorrelation-corrected, Student-t bands), checks the method
on each series' own history (hindcast), prices dated events with base-rate priors, and
scores whether a project's requirements are met at chosen horizons (default 5 and 10 years).

Subcommands
  cpi          Download US CPI-U (FRED CPIAUCSL) annual averages to cpi.csv
  forecast     Fan forecasts for every series in points.csv  -> forecasts.csv, series_fit.csv
  hindcast     Re-run the method from past origins on each series -> hindcast.csv, calibration.csv
  feasibility  Probability each requirement is met by each horizon -> feasibility.csv, project_score.csv
  workbook     Readable Excel workbook of the run (needs openpyxl)

Input CSVs use the same schema as PARADIGMA (points.csv, constraints.csv, noncost.csv),
so a PARADIGMA run folder can be reused directly. Requires pandas, numpy, matplotlib;
uses scipy for Student-t quantiles when present (normal fallback otherwise).
"""
import argparse
import io
import json
import math
import os
import re
import sys
import urllib.request

import numpy as np
import pandas as pd

try:
    from scipy import stats as _st
except Exception:
    _st = None

THETA = 0.63          # Farmer & Lafond (2016) pooled MA(1) autocorrelation for annual log-cost changes
MIN_DIFFS = 5         # minimum yearly changes needed to estimate drift and volatility
HORIZONS = (5, 10)

# Base-rate priors for dated events. JUDGMENT DEFAULTS, tunable; each row can be overridden in events.csv.
# p_on_time = P(happens by the scheduled date); p_eventual = P(happens within 3 years after it).
EVENT_PRIORS = {
    ("patent_expiry", "granted"): (0.90, 0.97),   # term extensions / SPCs / continuations shift some dates
    ("regulation", "enacted"): (0.75, 0.92),      # statutory dates do slip (e.g. EU AI Act 2026 omnibus delay)
    ("regulation", "proposed"): (0.30, 0.55),
    ("standard", "published"): (0.95, 0.98),
    ("standard", "draft"): (0.45, 0.75),
    ("capacity", "announced"): (0.25, 0.45),      # hydrogen ~11% of 2030 pipeline past FID (IEA 2025)
    ("capacity", "fid"): (0.60, 0.85),
    ("capacity", "under_construction"): (0.75, 0.92),
    ("launch", "announced"): (0.45, 0.70),
    ("launch", "dated_contract"): (0.65, 0.85),
    ("other", "any"): (0.50, 0.70),
}
# Readiness priors: median years per TRL step and log-sd of the TOTAL duration (steps are strongly
# correlated: slow early steps predict slow later ones). JUDGMENT DEFAULTS from thin evidence
# (GPS TRL3->9 11.7y; Ford product ~16y; Gross et al. 2018 invention->market median ~18y).
READINESS_PRIORS = {"software": (0.6, 0.6), "electronics": (1.3, 0.6), "hardware": (2.0, 0.6),
                    "energy": (2.5, 0.6), "infrastructure": (3.0, 0.6), "bio": (2.2, 0.7),
                    "medical": (2.5, 0.7), "materials": (2.5, 0.7)}
# Long-range widening. Forecasts of physical technologies are less accurate at long range than those of
# information technology (Fye et al. 2013; Kott & Perconti 2018), so for physical domains the spread is
# multiplied by 1 + WIDEN * max(0, h - 5) / 5 beyond 5 years. Widening (not pulling probabilities to 50/50)
# keeps impossible targets at 0 and already-met ones near 1. JUDGMENT DEFAULTS, tunable with --widen.
WIDEN = {"information": 0.0, "physical": 0.4}
INFO_DOMAINS = {"software", "electronics", "information", "compute", "ai"}


# ----------------------------------------------------------------- helpers
def parse_date(d):
    d = str(d).strip()
    m = re.fullmatch(r"(\d{4})-?Q([1-4])", d, re.I)
    if m:
        return int(m.group(1)) + (int(m.group(2)) - 0.5) / 4, int(m.group(1))
    m = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-(\d{1,2}))?", d)
    if m:
        return int(m.group(1)) + (int(m.group(2)) - 0.5) / 12, int(m.group(1))
    m = re.fullmatch(r"(\d{4})(?:\.0)?", d)
    if m:
        return int(m.group(1)) + 0.5, int(m.group(1))
    raise ValueError(f"bad date: {d!r}")


def t_ppf(q, df):
    if _st is not None:
        return float(_st.t.ppf(q, df))
    from statistics import NormalDist
    return NormalDist().inv_cdf(q) * math.sqrt(df / max(df - 2, 1)) if df > 2 else NormalDist().inv_cdf(q) * 2


def t_cdf(x, df):
    if _st is not None:
        return float(_st.t.cdf(x, df))
    return 0.5 * (1 + math.erf(x / math.sqrt(2)))


def theil_sen(x, y):
    sl = [(y[j] - y[i]) / (x[j] - x[i]) for i in range(len(x)) for j in range(i + 1, len(x)) if x[j] != x[i]]
    s = float(np.median(sl))
    return s, float(np.median(np.asarray(y) - s * np.asarray(x)))


def R(f):
    return pd.read_csv(f, dtype=str).fillna("") if f and os.path.exists(f) else pd.DataFrame()


def num(v, default=np.nan):
    try:
        v = str(v).replace(",", "").strip()
        return float(v) if v not in ("", "nan") else default
    except Exception:
        return default


# --------------------------------------------------------------------- cpi
def cmd_cpi(a):
    url = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=CPIAUCSL"
    df = pd.read_csv(io.StringIO(urllib.request.urlopen(url, timeout=30).read().decode()))
    df.columns = ["date", "cpi"]
    df["cpi"] = pd.to_numeric(df.cpi, errors="coerce")
    df["year"] = pd.to_datetime(df.date).dt.year
    g = df.dropna().groupby("year").agg(cpi=("cpi", "mean"), months=("cpi", "size")).reset_index()
    g["source"] = url
    g.to_csv(a.out, index=False)
    print(f"wrote {a.out}: {g.year.min()}-{g.year.max()}")


# ------------------------------------------------------------ data loading
def load_series(a):
    """points.csv -> {(cid, tier): DataFrame(year, median, n, sources)} in real USD where applicable."""
    pts = R(a.points)
    if pts.empty:
        sys.exit(f"no points in {a.points}")
    cpi = pd.read_csv(a.cpi) if a.cpi and os.path.exists(a.cpi) else None
    cmap, base_cpi, base_year = {}, None, None
    if cpi is not None:
        full = cpi[cpi.months >= 10] if "months" in cpi else cpi
        base_year = int(a.base_year or full.year.max())
        cmap = dict(zip(cpi.year.astype(int), cpi.cpi.astype(float)))
        base_cpi = cmap[base_year]
    rows, rej = [], []
    for _, r in pts.iterrows():
        why = None
        try:
            t, y = parse_date(r["date"])
            v = float(r["value"])
            if v <= 0:
                why = "value must be > 0"
        except Exception as e:
            why, y, v = str(e), None, None
        cur = str(r.get("currency", "")).upper().strip()
        basis = str(r.get("price_basis", "")).lower().strip()
        if not why and not str(r.get("source_url", "")).strip():
            why = "missing source_url"
        if not why and cur not in ("", "USD", "NA", "N/A"):
            why = f"currency {cur}: convert to USD first"
        if not why and a.end and y > a.end:
            why = f"after history end {a.end}"
        if not why and cur == "USD" and basis.startswith("nominal"):
            if base_cpi is None or y not in cmap:
                why = f"no CPI for {y} (run cpi)"
            else:
                v = v * base_cpi / cmap[y]
        elif not why and cur == "USD" and basis.startswith("real-") and base_cpi is not None:
            ry = int(basis.split("-")[1])
            if ry in cmap:
                v = v * base_cpi / cmap[ry]
        if why:
            rej.append({**r.to_dict(), "reject_reason": why})
            continue
        rows.append({"constraint_id": r["constraint_id"], "tier": r.get("tier", "") or "all",
                     "year": y, "value": v, "src": str(r["source_url"]).strip()})
    if rej:
        pd.DataFrame(rej).to_csv(os.path.join(a.outdir, "rejected_points.csv"), index=False)
    d = pd.DataFrame(rows)
    out = {}
    for k, g in d.groupby(["constraint_id", "tier"]):
        s = g.groupby("year").agg(median=("value", "median"), n=("value", "size"), vmin=("value", "min"),
                                  vmax=("value", "max"), sources=("src", "nunique")).reset_index().sort_values("year")
        s.attrs["n_sources"] = g.src.nunique()
        out[k] = s
    return out, base_year, len(rej), d


def fit_rw(years, logv, drift="endpoint", theta=THETA):
    """Random walk with drift on log values, tolerant of missing years.
    Returns dict(mu, K, m, y0, t0) or None. A gap of Δ years counts as Δ steps with variance K²Δ."""
    years = np.asarray(years, float)
    logv = np.asarray(logv, float)
    if len(years) < 2:
        return None
    dt = np.diff(years)
    dy = np.diff(logv)
    m = int(round(dt.sum()))                     # number of yearly steps spanned
    if drift == "endpoint":
        mu = (logv[-1] - logv[0]) / (years[-1] - years[0])
    else:
        mu, _ = theil_sen(years, logv)
    z = (dy - mu * dt) / np.sqrt(dt)
    nd = len(z)
    if nd < 2:
        return None
    K = float(np.sqrt(np.sum((z - z.mean()) ** 2) / (nd - 1)))
    if drift == "robust":
        # outlier-resistant volatility (1.4826 x median absolute deviation); used only when it wins the backtest
        K = float(1.4826 * np.median(np.abs(z - np.median(z)))) or K
    # forecast from the last observed (yearly-median) value, as in Farmer & Lafond. Smoothing the anchor
    # or using a robust drift made hindcast bands too narrow in testing, so the published form is the default.
    anchor = logv[-1]
    return {"mu": float(mu), "K": K, "m": m, "n_diffs": nd, "y0": float(anchor), "t0": float(years[-1]),
            "theta": theta}


def fan(fit, tau, q, df=None):
    """Quantile q of log value tau years after t0 (Farmer & Lafond 2016, eq. for A*)."""
    m, th = max(fit["m"], 2), fit["theta"]
    A = -2 * th + (1 + 2 * (m - 1) * th / m + th ** 2) * (tau + tau ** 2 / m)
    sd = fit["K"] * math.sqrt(max(A, 1e-12) / (1 + th ** 2))
    df = df or max(fit["n_diffs"] - 1, 1)
    return fit["y0"] + fit["mu"] * tau + t_ppf(q, df) * sd, sd


def p_better_than(fit, tau, target, better, floor=np.nan, ceiling=np.nan, widen=1.0):
    """P(value at t0+tau is at least as good as target), respecting hard bounds."""
    lt = math.log(target)
    if better == "lower" and not np.isnan(floor) and target < floor:
        return 0.0
    if better == "higher" and not np.isnan(ceiling) and target > ceiling:
        return 0.0
    mean, sd = fan(fit, tau, 0.5)
    df = max(fit["n_diffs"] - 1, 1)
    zc = (lt - mean) / max(sd * widen, 1e-12)
    return t_cdf(zc, df) if better == "lower" else 1 - t_cdf(zc, df)


def constraints_meta(path):
    c = R(path)
    meta = {}
    for _, r in c.iterrows():
        meta[r["constraint_id"]] = {
            "name": r.get("name", ""), "unit": r.get("unit", ""), "better": (r.get("better", "") or "lower").lower(),
            "floor": num(r.get("floor", "")), "ceiling": num(r.get("ceiling", "")),
            "bound_source": r.get("bound_source", ""), "domain": (r.get("domain", "") or "").lower()}
    return meta


def apply_scenarios(fit, cid, tier, scen):
    """Yield (scenario_name, fit) — the trend fit plus any sourced overrides from scenarios.csv."""
    yield "trend", fit
    if scen.empty:
        return
    s = scen[(scen.constraint_id == cid) & (scen.get("tier", pd.Series([""] * len(scen))).isin(["", tier]))]
    for _, r in s.iterrows():
        f = dict(fit)
        ch = num(r.get("annual_change_pct", ""))
        if not np.isnan(ch):
            f["mu"] = math.log(1 + ch / 100)
        km = num(r.get("volatility_multiplier", ""))
        if not np.isnan(km):
            f["K"] = fit["K"] * km
        yield r["scenario"], f



# --------------------------------------------------------------- auto-fixes
# Three problems are repaired automatically (switch off with --no-autofix):
#  1. mixed products in one series   -> track the best available value each year (price-performance frontier)
#  2. the trend changed               -> statistical test of the drift in the last 15/10/8 years vs the years
#                                        before; if significant, fit on the recent window
#  2b. outliers                       -> robust trend (Theil-Sen) and volatility (MAD) when a yearly change is extreme
#  3. the trend runs into a limit     -> model the remaining gap to the physical floor/ceiling, which slows
#                                        progress as it approaches the limit instead of hitting it and stopping
# Fixes are chosen by explicit tests, not by picking the best of many backtests (that overfits noise).
# A fourth (too little data -> uninformative ranges) cannot be fixed by arithmetic; the skill sends a research
# agent for denser data instead.
WINDOWS = (15, 10, 8)
BREAK_Z = 3.2          # trend-change test threshold (strict: 3 windows are tried; ~1-2% false alarms on clean data)
OUTLIER_Z = 5.5        # a yearly change this many robust SDs from typical counts as an outlier


def yearly(g, agg, better):
    f = "median" if agg == "median" else ("min" if better == "lower" else "max")
    s = g.groupby("year").agg(median=("value", f), n=("value", "size"), vmin=("value", "min"),
                              vmax=("value", "max"), sources=("src", "nunique")).reset_index().sort_values("year")
    if agg == "frontier":
        # best available so far: an older cheaper/better product stays buyable, so the frontier never moves back
        s["median"] = s["median"].cummin() if better == "lower" else s["median"].cummax()
    s.attrs["n_sources"] = g.src.nunique()
    return s


def make_model(years, vals, better, bound, transform, drift):
    years, vals = np.asarray(years, float), np.asarray(vals, float)
    if transform == "gap":
        gap = (bound - vals) if better == "higher" else (vals - bound)
        if len(gap) == 0 or np.any(gap <= 0):
            return None
        y = np.log(gap)
    else:
        y = np.log(vals)
    f = fit_rw(years, y, drift)
    if f is None:
        return None
    f.update(transform=transform, bound=bound, better=better)
    return f


def m_value(m, tau, q, widen=1.0):
    """Quantile q of the VALUE tau years after the model's last data year."""
    mean, sd = fan(m, tau, 0.5)
    df = max(m["n_diffs"] - 1, 1)
    if m.get("transform") == "gap":
        qq = 1 - q if m["better"] == "higher" else q
        g = math.exp(mean + t_ppf(qq, df) * sd * widen)
        return max(m["bound"] - g, 1e-300) if m["better"] == "higher" else m["bound"] + g
    return math.exp(mean + t_ppf(q, df) * sd * widen)


def m_prob(m, tau, target, floor=np.nan, ceiling=np.nan, widen=1.0):
    """P(value at t0+tau is at least as good as target)."""
    if m.get("transform") == "gap":
        b = m["bound"]
        gap_t = (b - target) if m["better"] == "higher" else (target - b)
        if gap_t <= 0:
            return 0.0
        mean, sd = fan(m, tau, 0.5)
        return t_cdf((math.log(gap_t) - mean) / max(sd * widen, 1e-12), max(m["n_diffs"] - 1, 1))
    return p_better_than(m, tau, target, m["better"], floor, ceiling, widen)


def crossing_year(m, target):
    """Year the median path reaches target, 'already met', or 'trend moves away'."""
    if m.get("transform") == "gap":
        b = m["bound"]
        gap_t = (b - target) if m["better"] == "higher" else (target - b)
        if gap_t <= 0:
            return "never: beyond floor/ceiling"
        need, met_if_pos = math.log(gap_t) - m["y0"], True
    else:
        need, met_if_pos = math.log(target) - m["y0"], m["better"] == "lower"
    if (need >= 0) == met_if_pos or need == 0:
        return "already met"
    if m["mu"] != 0 and np.sign(need) == np.sign(m["mu"]):
        return round(m["t0"] + need / m["mu"], 1)
    return "trend moves away"


def backtest(years, vals, better, bound, transform, window, drift, min_diffs, last_k=12, max_h=5):
    """Out-of-sample check over the last `last_k` origins: {(origin, year): (abs log error, inside 80%)}."""
    years, vals = np.asarray(years, float), np.asarray(vals, float)
    last, res = years.max(), {}
    for o in years:
        if o < last - last_k or o >= last:
            continue
        mask = (years <= o) & ((years > o - window) if window else True)
        m = make_model(years[mask], vals[mask], better, bound, transform, drift)
        if m is None or m["n_diffs"] < min_diffs:
            continue
        for j in np.where((years > o) & (years <= o + max_h))[0]:
            tau = years[j] - m["t0"]
            med, lo, hi = m_value(m, tau, .5), m_value(m, tau, .1), m_value(m, tau, .9)
            res[(int(o), int(years[j]))] = (abs(math.log(vals[j]) - math.log(med)), lo <= vals[j] <= hi)
    return res


def choose_model(g, mt, a, horizon_end):
    """Pick aggregation, window and transform for one series. Returns (model, yearly_series, info) or (None, s, info)."""
    better = mt.get("better", "lower") or "lower"
    bound = mt.get("floor", np.nan) if better == "lower" else mt.get("ceiling", np.nan)
    s_med = yearly(g, "median", better)
    multi = s_med[s_med.n >= 2]
    spread = float(np.median(np.log(multi.vmax / multi.vmin))) if len(multi) >= 2 else 0.0
    mixed = spread > math.log(3)
    autofix = not getattr(a, "no_autofix", False)
    agg = "frontier" if (mixed and autofix) else "median"
    s = yearly(g, agg, better)
    info = {"aggregation": agg, "window": "all", "transform": "log", "fixes": [], "spread": spread, "mixed": mixed}
    if mixed and autofix:
        info["fixes"].append("mixed products: switched to the best available value per year (frontier)")
    yrs, vals = s.year.values.astype(float), s["median"].values.astype(float)
    base = make_model(yrs, vals, better, bound, "log", a.drift)
    if base is None or base["n_diffs"] < a.min_diffs:
        return None, s, info
    if not autofix:
        return base, s, info
    bound_ok = not np.isnan(bound) and bool(np.all(((bound - vals) if better == "higher" else (vals - bound)) > 0))
    drift = a.drift
    # Fix: outliers. A yearly change more than OUTLIER_Z robust SDs from the typical change -> robust trend + volatility.
    lv = np.log(vals)
    z_inc = np.diff(lv) / np.sqrt(np.diff(yrs))
    mad = 1.4826 * np.median(np.abs(z_inc - np.median(z_inc)))
    if mad > 0 and len(z_inc) >= 6:
        worst = float(np.max(np.abs(z_inc - np.median(z_inc))) / mad)
        if worst > OUTLIER_Z:
            drift = "robust"
            info["fixes"].append(f"outlier: one yearly change is {worst:.1f} robust SDs from typical; used a robust "
                                 f"trend and volatility")
    # Fix: the trend changed. Test drift in the last w years against the years before (Farmer-Lafond MA(1) errors).
    chosen_w, best_z = None, 0.0
    for w in WINDOWS:
        rec, old = yrs > yrs.max() - w, yrs <= yrs.max() - w
        if rec.sum() < a.min_diffs + 1 or old.sum() < a.min_diffs + 1:
            continue
        f1 = make_model(yrs[old], vals[old], better, bound, "log", drift)
        f2 = make_model(yrs[rec], vals[rec], better, bound, "log", drift)
        if f1 is None or f2 is None:
            continue
        def var_mu(f):
            n, th = max(f["m"], 1), THETA
            return f["K"] ** 2 * ((n - 1) * (1 + th) ** 2 + 1 + th ** 2) / ((1 + th ** 2) * n ** 2)
        z = abs(f2["mu"] - f1["mu"]) / math.sqrt(var_mu(f1) + var_mu(f2))
        if z > BREAK_Z and z > best_z:
            chosen_w, best_z, rates = w, z, (f1["mu"], f2["mu"])
    tr = "log"
    # Fix: a log trend that crosses the physical limit inside the horizon predicts impossible values -> gap model
    sel = yrs > yrs.max() - chosen_w if chosen_w else np.ones(len(yrs), bool)
    mb = make_model(yrs[sel], vals[sel], better, bound, "log", drift)
    hits = False
    if bound_ok and mb is not None and mb["mu"] != 0 and np.sign(math.log(bound) - mb["y0"]) == np.sign(mb["mu"]):
        hits = mb["t0"] + (math.log(bound) - mb["y0"]) / mb["mu"] <= horizon_end
        tr = "gap" if hits else "log"
    m = make_model(yrs[sel], vals[sel], better, bound, tr, drift)
    if m is None or m["n_diffs"] < a.min_diffs:
        return base, s, info
    info.update(window=f"last {chosen_w}y" if chosen_w else "all", transform=tr, drift=drift)
    if chosen_w:
        info["fixes"].append(f"trend changed: last {chosen_w} years moved {(math.exp(rates[1]) - 1) * 100:+.1f}%/yr vs "
                             f"{(math.exp(rates[0]) - 1) * 100:+.1f}%/yr before (z = {best_z:.1f}); fitted on the last "
                             f"{chosen_w} years")
    if tr == "gap":
        info["fixes"].append("approaching a physical limit: the straight trend would cross it; modelled the remaining "
                             "gap to the limit so progress slows near it (an assumption from physics: the history "
                             "itself does not show the slowdown yet)")
    # report, but do not select on, the backtest of the chosen model vs the default (information only)
    bt_def = backtest(yrs, vals, better, bound, "log", None, a.drift, a.min_diffs)
    bt_new = backtest(yrs, vals, better, bound, tr, chosen_w, drift, a.min_diffs)
    k = set(bt_def) & set(bt_new)
    if len(k) >= 4 and info["fixes"]:
        e0, e1 = np.mean([bt_def[x][0] for x in k]), np.mean([bt_new[x][0] for x in k])
        info["backtest_gain"] = round(1 - e1 / e0, 3) if e0 > 0 else 0.0
    return m, s, info


# ---------------------------------------------------------------- forecast
def cmd_forecast(a):
    series, base_year, nrej, raw = load_series(a)
    meta = constraints_meta(a.constraints)
    scen = R(a.scenarios)
    horizons = sorted(set(int(h) for h in a.horizons.split(",")))
    origin = a.end
    hz_end = origin + max(horizons)
    calf = R(os.path.join(a.outdir, "calibration.csv"))
    cal = {(r.constraint_id, r.tier): r.to_dict() for _, r in calf.iterrows()} if len(calf) else {}
    qs = [0.025, 0.10, 0.25, 0.5, 0.75, 0.90, 0.975]
    fits, fc, used, models = [], [], {}, {}
    for (cid, tier), g in raw.groupby(["constraint_id", "tier"]):
        mt = meta.get(cid, {"better": "lower", "floor": np.nan, "ceiling": np.nan, "name": cid, "unit": ""})
        fit, s, info = choose_model(g, mt, a, hz_end)
        used[(cid, tier)] = s
        row = {"constraint_id": cid, "tier": tier, "first_year": int(s.year.min()), "last_year": int(s.year.max()),
               "n_years": len(s), "n_points": int(s.n.sum()), "n_sources": s.attrs["n_sources"]}
        if fit is None:
            row.update({"status": f"insufficient (<{a.min_diffs} yearly changes)", "grade": "D"})
            fits.append(row)
            continue
        models[(cid, tier)] = fit
        stale = origin - int(s.year.max())
        brk = ""
        if len(s) >= 8 and info["window"] == "all":
            half = len(s) // 2
            f1 = fit_rw(s.year.values[:half + 1], np.log(s["median"].values[:half + 1]), info.get("drift", a.drift))
            f2 = fit_rw(s.year.values[half:], np.log(s["median"].values[half:]), info.get("drift", a.drift))
            if f1 and f2:
                se = math.sqrt(f1["K"] ** 2 / max(f1["m"], 1) + f2["K"] ** 2 / max(f2["m"], 1))
                if abs(f2["mu"] - f1["mu"]) > 2 * se:
                    brk = "trend changed between halves"
        g_ = ("A" if fit["n_diffs"] >= 15 and row["n_sources"] >= 3 and stale <= 2 else
              "B" if fit["n_diffs"] >= 8 and row["n_sources"] >= 2 and stale <= 3 else "C")
        cv = cal.get((cid, tier))
        hv = (cv or {}).get("verdict", "not run")
        unfixed = []
        if info["mixed"] and info["aggregation"] == "median":
            unfixed.append("mixed products")
        if str(hv).startswith(("biased", "bands too narrow")) and info["window"] == "all":
            unfixed.append("hindcast biased")
        v_now = m_value(fit, 0.001, 0.5)
        v_end = m_value(fit, hz_end - fit["t0"], 0.5)
        span = hz_end - fit["t0"]
        chg = (v_end / v_now) ** (1 / span) - 1 if span > 0 else 0
        row.update({"status": "ok", "aggregation": info["aggregation"], "fit_window": info["window"],
                    "model": "log trend" if fit["transform"] == "log" else "gap to limit",
                    "fixes_applied": " | ".join(info["fixes"]) or "none needed",
                    "drift_method": info.get("drift", a.drift),
                    "annual_change_pct": round((math.exp(fit["mu"]) - 1) * 100, 2) if fit["transform"] == "log"
                    else round(chg * 100, 2),
                    "volatility_K": round(fit["K"], 4), "m_years": fit["m"], "theta": fit["theta"],
                    "years_since_last_data": stale, "break_check": brk or "none detected",
                    "within_year_spread_x": round(math.exp(info["spread"]), 2),
                    "hindcast_verdict_default_model": hv,
                    "floor": mt["floor"], "ceiling": mt["ceiling"]})
        lo, hi = m_value(fit, span, 0.10), m_value(fit, span, 0.90)
        width = max(hi, lo) / max(min(hi, lo), 1e-300)
        row["range80_width_x_at_max_horizon"] = round(width, 1)
        row["informative"] = ("yes" if width <= 10 else "weak (80% range > 10x)" if width <= 100 else
                              "NO: 80% range > 100x even after fixes; needs denser data")
        if width > 100:
            unfixed.append("uninformative")
        if brk:
            unfixed.append("trend changed")
        row["unresolved_warnings"] = ", ".join(unfixed)
        if unfixed and g_ in ("A", "B"):
            g_ = "C"
        if info["fixes"] and g_ == "A":
            g_ = "B"          # a repaired series is never top grade
        row["grade"] = g_
        fits.append(row)
        for sname, f in apply_scenarios(fit, cid, tier, scen):
            for yr in range(int(f["t0"]) + 1, hz_end + 1):
                tau = yr - f["t0"]
                rec = {"constraint_id": cid, "tier": tier, "scenario": sname, "year": yr, "tau": tau,
                       "horizon": yr - origin}
                for q in qs:
                    v = m_value(f, tau, q)
                    if mt["better"] == "lower" and not np.isnan(mt["floor"]) and v < mt["floor"]:
                        v, rec["bound_hit"] = mt["floor"], True
                    if mt["better"] == "higher" and not np.isnan(mt["ceiling"]) and v > mt["ceiling"]:
                        v, rec["bound_hit"] = mt["ceiling"], True
                    rec[f"q{int(q * 1000):03d}"] = v
                fc.append(rec)
    fdf, fcdf = pd.DataFrame(fits), pd.DataFrame(fc)
    fdf.to_csv(os.path.join(a.outdir, "series_fit.csv"), index=False)
    fcdf.to_csv(os.path.join(a.outdir, "forecasts.csv"), index=False)
    json.dump({"origin_year": origin, "base_year": base_year, "theta": THETA, "drift": a.drift, "autofix": not a.no_autofix,
               "horizons": horizons, "rejected_points": nrej},
              open(os.path.join(a.outdir, "forecast_meta.json"), "w"), indent=2)
    cols = ["constraint_id", "tier", "status", "grade", "aggregation", "fit_window", "model", "annual_change_pct",
            "volatility_K", "informative", "unresolved_warnings"]
    print(fdf[[c for c in cols if c in fdf]].to_string(index=False))
    print(f"\norigin {origin}, horizons {horizons}, rejected points {nrej}, dollars of {base_year}")
    if a.plots:
        plot_fans(used, fdf, fcdf, meta, a, base_year, horizons)


# ---------------------------------------------------------------- hindcast
def cmd_hindcast(a):
    """For each series with enough history, forecast from every past origin and score against what happened."""
    series, *_ = load_series(a)
    rows = []
    for (cid, tier), s in series.items():
        yrs, lv = s.year.values, np.log(s["median"].values)
        for i in range(len(yrs)):
            sub_y, sub_v = yrs[:i + 1], lv[:i + 1]
            f = fit_rw(sub_y, sub_v, a.drift)
            if f is None or f["n_diffs"] < a.min_diffs:
                continue
            for j in range(i + 1, len(yrs)):
                tau = yrs[j] - yrs[i]
                if tau > a.max_tau:
                    break
                mean, sd = fan(f, tau, 0.5)
                df = max(f["n_diffs"] - 1, 1)
                z = (lv[j] - mean) / sd
                rows.append({"constraint_id": cid, "tier": tier, "origin": int(yrs[i]), "target_year": int(yrs[j]),
                             "tau": int(tau), "log_error": round(lv[j] - mean, 4), "z": round(z, 3),
                             "in80": abs(z) <= t_ppf(0.90, df), "in95": abs(z) <= t_ppf(0.975, df)})
    h = pd.DataFrame(rows)
    h.to_csv(os.path.join(a.outdir, "hindcast.csv"), index=False)
    if h.empty:
        print("no series long enough for a hindcast")
        return
    cal = h.groupby(["constraint_id", "tier"]).agg(n=("z", "size"), cover80=("in80", "mean"), cover95=("in95", "mean"),
                                                    bias_log=("log_error", "mean")).reset_index()
    pooled = pd.DataFrame([{"constraint_id": "ALL", "tier": "pooled", "n": len(h), "cover80": h.in80.mean(),
                            "cover95": h.in95.mean(), "bias_log": h.log_error.mean()}])
    cal = pd.concat([cal, pooled], ignore_index=True)
    meta = constraints_meta(a.constraints)
    def _v(r):
        if r.n < 10:
            return "too few checks"
        if abs(r.bias_log) > 0.3:
            better = meta.get(r.constraint_id, {}).get("better", "lower")
            optimistic = (r.bias_log > 0) == (better == "lower")
            return (f"biased: past forecasts were too {'optimistic' if optimistic else 'pessimistic'} "
                    f"(x{math.exp(abs(r.bias_log)):.2f}); the trend changed, refit on recent years")
        if r.cover80 < 0.65:
            return "bands too narrow: widen (volatility_multiplier) or distrust"
        if r.cover80 > 0.95:
            return "bands wide (conservative)"
        return "calibrated"
    cal["verdict"] = cal.apply(_v, axis=1)
    cal.loc[cal.constraint_id == "ALL", "verdict"] = cal.loc[cal.constraint_id == "ALL", "verdict"].str.replace(
        "the trend changed, refit on recent years", "check each series")
    for c in ("cover80", "cover95", "bias_log"):
        cal[c] = cal[c].round(3)
    cal.to_csv(os.path.join(a.outdir, "calibration.csv"), index=False)
    print(cal.to_string(index=False))


# ------------------------------------------------------------- feasibility
def event_prob(ev, year):
    """P(event has happened by `year`)."""
    try:
        when = parse_date(ev["date_scheduled"])[0]
    except Exception:
        return np.nan
    key = (ev.get("type", "other").lower(), ev.get("status", "any").lower())
    p_on, p_ev = EVENT_PRIORS.get(key, EVENT_PRIORS[("other", "any")])
    if num(ev.get("p_on_time_override", "")) == num(ev.get("p_on_time_override", "")):
        p_on = num(ev["p_on_time_override"])
    if num(ev.get("p_eventual_override", "")) == num(ev.get("p_eventual_override", "")):
        p_ev = num(ev["p_eventual_override"])
    t = year + 0.999
    if t < when:
        return 0.0
    if t >= when + 3:
        return p_ev
    return p_on + (p_ev - p_on) * (t - when) / 3


def readiness_prob(current, needed, domain, years, override_step=np.nan, widen=1.0):
    steps = max(0.0, needed - current)
    if steps == 0:
        return 1.0
    med_step, ls = READINESS_PRIORS.get(domain, READINESS_PRIORS["hardware"])
    if not np.isnan(override_step):
        med_step = override_step
    median_total = med_step * steps
    if years <= 0:
        return 0.0
    z = (math.log(years) - math.log(median_total)) / (ls * widen)
    return 0.5 * (1 + math.erf(z / math.sqrt(2)))


def widen_factor(domain, h, physical=None):
    w = WIDEN["information"] if domain in INFO_DOMAINS else (WIDEN["physical"] if physical is None else physical)
    return 1 + w * max(0.0, h - 5) / 5


def cmd_feasibility(a):
    req = R(a.requirements)
    if req.empty:
        sys.exit(f"no requirements in {a.requirements}")
    horizons = sorted(set(int(h) for h in a.horizons.split(",")))
    origin = a.end
    meta = constraints_meta(a.constraints)
    _s, _b, _n, raw = load_series(a) if os.path.exists(a.points) else ({}, None, 0, pd.DataFrame())
    scen = R(a.scenarios)
    events = R(a.events)
    evmap = {r["event_id"]: r for _, r in events.iterrows()} if len(events) else {}
    rows = []
    all_h = sorted(set(horizons) | set(range(0, max(horizons) + 1)))
    for _, r in req.iterrows():
        kind = r["kind"].strip().lower()
        dom = (r.get("domain", "") or meta.get(r.get("constraint_id", ""), {}).get("domain", "") or a.domain).lower()
        base = {"req_id": r["req_id"], "requirement": r.get("description", ""), "kind": kind, "domain": dom,
                "critical": str(r.get("critical", "yes")).lower() in ("yes", "y", "1", "true", "")}
        probs, raw_p = {}, {}
        note = ""
        if kind == "curve":
            cid, tier = r["constraint_id"], r.get("tier", "") or "all"
            mt = dict(meta.get(cid, {"better": "lower", "floor": np.nan, "ceiling": np.nan}))
            better = (r.get("better", "") or mt["better"]).lower()
            mt["better"] = better
            target = num(r["target_value"])
            g = raw[(raw.constraint_id == cid) & (raw.tier == tier)] if len(raw) else raw
            f, s, info = choose_model(g, mt, a, origin + max(horizons)) if len(g) else (None, None, {})
            if f is None:
                note = "no usable curve; treat as judgment"
            else:
                sc = r.get("scenario", "") or "trend"
                fits = dict(apply_scenarios(f, cid, tier, scen))
                f = fits.get(sc, f)
                base["current_value"] = round(float(s["median"].values[-1]), 6)
                base["target_value"] = target
                for h in all_h:
                    tau = max(origin + h - f["t0"], 0.01)
                    raw_p[h] = m_prob(f, tau, target, mt["floor"], mt["ceiling"])
                    probs[h] = m_prob(f, tau, target, mt["floor"], mt["ceiling"], widen_factor(dom, h, a.widen))
                base["median_crossing_year"] = crossing_year(f, target)
                note = (f"scenario={sc}; {cid}/{tier}; {info.get('aggregation')}, window {info.get('window')}, "
                        f"{'gap to limit' if f.get('transform') == 'gap' else 'log trend'}"
                        + (f"; fixes: {' | '.join(info['fixes'])}" if info.get("fixes") else ""))
        elif kind == "event":
            ev = evmap.get(r["event_id"])
            if ev is None:
                note = f"event {r['event_id']} missing from events.csv"
            else:
                for h in all_h:
                    probs[h] = event_prob(ev, origin + h)
                note = f"{ev.get('type','')}/{ev.get('status','')} due {ev.get('date_scheduled','')}"
        elif kind == "readiness":
            cur, need = num(r["current_trl"]), num(r.get("needed_trl", 9) or 9)
            for h in all_h:
                raw_p[h] = readiness_prob(cur, need, dom, h + a.since_offset, num(r.get("years_per_step", "")))
                probs[h] = readiness_prob(cur, need, dom, h + a.since_offset, num(r.get("years_per_step", "")),
                                          widen_factor(dom, h, a.widen))
            note = f"TRL {cur:g}->{need:g}, prior {dom or 'hardware'}"
        elif kind == "judgment":
            pts = {int(k[2:-1]): num(r[k]) for k in r.index if re.fullmatch(r"p_\d+y", k) and num(r[k]) == num(r[k])}
            if pts:
                xs = sorted(pts)
                for h in all_h:
                    probs[h] = float(np.interp(h, [0] + xs, [num(r.get("p_now", 0) or 0)] + [pts[x] for x in xs]))
            note = f"source: {r.get('source_url','')}"
        for h in horizons:
            p, pr = probs.get(h, np.nan), raw_p.get(h, probs.get(h, np.nan))
            base[f"p_model_{h}y"] = round(pr, 3) if pr == pr else np.nan
            base[f"p_{h}y"] = round(p, 3) if p == p else np.nan
        base["path"] = json.dumps({h: round(probs[h], 3) for h in all_h if h in probs})
        base["note"] = note
        rows.append(base)
    fz = pd.DataFrame(rows)
    fz.to_csv(os.path.join(a.outdir, "feasibility.csv"), index=False)

    # project score
    crit = fz[fz.critical]
    out = []
    for h in horizons:
        ps = crit[f"p_{h}y"].astype(float)
        if ps.isna().any():
            out.append({"horizon_years": h, "year": origin + h, "status": "incomplete: some requirements unscored"})
            continue
        joint = float(np.prod(ps))
        weak = crit.loc[ps.idxmin()]
        out.append({"horizon_years": h, "year": origin + h, "joint_if_independent": round(joint, 3),
                    "weakest_link_bound": round(float(ps.min()), 3), "weakest_requirement": weak.req_id,
                    "verdict": verdict(joint, float(ps.min()))})
    # start window: latest year to start so that, after lead time, P(all critical met) >= 0.5
    lead = a.lead_years
    paths = [json.loads(p) for p in crit.path]
    if paths and all(paths):
        js = {}
        for h in range(0, max(horizons) + 1):
            vals = [p.get(str(h), p.get(h)) for p in paths]
            if all(v is not None for v in vals):
                js[h] = float(np.prod(vals))
        ready = [h for h, v in js.items() if v >= 0.5]
        out.append({"horizon_years": "start-window", "year": (origin + ready[0] - lead) if ready else "",
                    "status": (f"P(all critical met) first >= 0.5 in {origin + ready[0]}; start building by "
                               f"{origin + ready[0] - lead} with a {lead}-year lead" if ready else
                               f"P(all critical met) stays < 0.5 through {origin + max(js) if js else origin}"),
                    "joint_path": json.dumps({origin + h: round(v, 3) for h, v in js.items()})})
    ps_df = pd.DataFrame(out)
    ps_df.to_csv(os.path.join(a.outdir, "project_score.csv"), index=False)
    print(fz.drop(columns=["path"]).to_string(index=False))
    print()
    print(ps_df.to_string(index=False))
    if a.plots:
        plot_feasibility(fz, ps_df, horizons, origin, a)


def verdict(joint, weakest):
    if weakest >= 0.8 and joint >= 0.6:
        return "LIKELY FEASIBLE"
    if joint >= 0.35:
        return "PLAUSIBLE"
    if joint >= 0.12:
        return "LONG SHOT"
    return "UNLIKELY"


# ------------------------------------------------------------------- plots
# One look for every chart: white page, pale-blue plot area, near-black ink, one blue accent
# (PARADIGMA uses the same layout in pink, so the two skills read as a pair).
INK, MUTED, DEEP, SKY, MIST, WASH, GRIDC = "#111111", "#55555E", "#2F6FB5", "#8DB8E3", "#C9DDF2", "#F4F8FD", "#D6E4F2"


def _mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 10, "axes.edgecolor": "#BBBBC2", "axes.labelcolor": INK,
                         "xtick.color": MUTED, "ytick.color": MUTED})
    return plt


def plot_fans(series, fdf, fcdf, meta, a, base_year, horizons):
    plt = _mpl()
    from matplotlib.ticker import FuncFormatter, LogLocator, NullFormatter, MaxNLocator
    os.makedirs(os.path.join(a.outdir, "plots"), exist_ok=True)
    req = R(a.requirements) if a.requirements else pd.DataFrame()
    for _, fr in fdf[fdf.status == "ok"].iterrows():
        cid, tier = fr.constraint_id, fr.tier
        s = series[(cid, tier)]
        c = fcdf[(fcdf.constraint_id == cid) & (fcdf.tier == tier) & (fcdf.scenario == "trend")]
        fig, ax = plt.subplots(figsize=(8.4, 4.8), facecolor="white")
        ax.set_facecolor(WASH)
        ax.grid(color=GRIDC, lw=0.8)
        lab = "history (best value per year)" if fr.get("aggregation") == "frontier" else "history (yearly median)"
        ax.plot(s.year, s["median"], "o-", color=INK, ms=4, lw=1.4, mfc="white", label=lab)
        wl = str(fr.get("fit_window", "all"))
        if wl.startswith("last"):
            w0 = s.year.max() - int(wl.split()[1].rstrip("y"))
            ax.axvspan(w0 + 0.5, s.year.max(), color=MIST, alpha=0.35, lw=0, label=f"years used for the trend ({wl})")
        x = np.r_[s.year.values[-1], c.year.values]
        last = s["median"].values[-1]
        for lo, hi, col, lab in (("q025", "q975", MIST, "95% range"), ("q100", "q900", SKY, "80% range"),
                                 ("q250", "q750", DEEP, "50% range")):
            ax.fill_between(x, np.r_[last, c[lo].values], np.r_[last, c[hi].values], color=col,
                            alpha=0.55 if col != DEEP else 0.35, lw=0, label=lab)
        ax.plot(x, np.r_[last, c.q500.values], "--", color=INK, lw=1.4, label="median forecast")
        for o, sc in fcdf[(fcdf.constraint_id == cid) & (fcdf.tier == tier) & (fcdf.scenario != "trend")].groupby("scenario"):
            ax.plot(sc.year, sc.q500, ":", lw=1.6, color=MUTED, label=f"scenario: {o}")
        mt = meta.get(cid, {})
        if mt and not np.isnan(mt.get("floor", np.nan)):
            ax.axhline(mt["floor"], color=MUTED, lw=1, ls="-.", label="floor")
        if mt and not np.isnan(mt.get("ceiling", np.nan)):
            ax.axhline(mt["ceiling"], color=MUTED, lw=1, ls="-.", label="ceiling (physical limit)")
        if len(req) and "constraint_id" in req:
            for _, rq in req[(req.constraint_id == cid)].iterrows():
                tv = num(rq.get("target_value", ""))
                if tv == tv:
                    ax.axhline(tv, color=DEEP, lw=1.6, label=f"target {rq.req_id}")
        for h in horizons:
            ax.axvline(a.end + h, color=GRIDC, lw=1.2)
            ax.annotate(f"+{h}y", (a.end + h, 0), xycoords=("data", "axes fraction"), xytext=(0, -26),
                        textcoords="offset points", ha="center", color=DEEP, fontsize=8, fontweight="bold")
        ax.set_yscale("log")
        lo_, hi_ = ax.get_ylim()
        decades = math.log10(hi_ / lo_)
        subs = (1.0, 2.0, 5.0) if decades <= 4 else (1.0,)
        ax.yaxis.set_major_locator(LogLocator(base=10, subs=subs, numticks=12 if decades <= 12 else 8))
        ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.6g}" if 1e-4 <= v < 1e7 else f"{v:.0e}"))
        ax.yaxis.set_minor_formatter(NullFormatter())
        ax.xaxis.set_major_locator(MaxNLocator(integer=True))
        unit = mt.get("unit", "") if mt else ""
        money = "USD" in str(unit).upper() or "$" in str(unit)
        ax.set_ylabel(f"{unit} ({base_year} dollars)" if base_year and money else unit)
        for sp in ("top", "right"):
            ax.spines[sp].set_visible(False)
        ax.legend(fontsize=7.8, frameon=False, loc="center left", bbox_to_anchor=(1.01, 0.5))
        fig.text(0.01, 0.97, f"{cid} {mt.get('name', cid) if mt else cid}", fontsize=12.5, fontweight="bold", color=INK, va="top")
        fixes = str(fr.get("fixes_applied", ""))
        how = f"{fr.get('aggregation', 'median')} per year · fit {fr.get('fit_window', 'all')} · {fr.get('model', 'log trend')}"
        unres = str(fr.get("unresolved_warnings", ""))
        unres = "" if unres in ("", "nan") else f" · STILL: {unres.upper()}"
        fig.text(0.01, 0.905, f"series {tier} · {how} · volatility {fr.volatility_K:.2f} · grade {fr.grade}{unres}",
                 fontsize=8.6, color=DEEP if unres else MUTED, va="top")
        if fixes not in ("", "nan", "none needed"):
            import textwrap
            fig.text(0.01, 0.865, "auto-fixed: " + "\n".join(textwrap.wrap(fixes, 150)), fontsize=7.8, color=DEEP, va="top")
        fig.subplots_adjust(left=0.1, right=0.74, top=0.80, bottom=0.1)
        fig.savefig(os.path.join(a.outdir, "plots", f"fan_{cid}_{tier}.png"), dpi=150, bbox_inches="tight",
                    pad_inches=0.25, facecolor="white")
        plt.close(fig)


def plot_feasibility(fz, ps_df, horizons, origin, a):
    plt = _mpl()
    os.makedirs(os.path.join(a.outdir, "plots"), exist_ok=True)
    fig, ax = plt.subplots(figsize=(9, 0.55 * len(fz) + 2.4), facecolor="white")
    ax.set_facecolor(WASH)
    y = np.arange(len(fz))[::-1]
    w = 0.8 / len(horizons)
    cols = [SKY, DEEP, INK]
    for i, h in enumerate(horizons):
        v = fz[f"p_{h}y"].astype(float).fillna(0).values
        yy_ = y + ((len(horizons) - 1) / 2 - i) * w
        ax.barh(yy_, v, height=w, color=cols[i % 3], label=f"by {origin + h} (+{h}y)")
        for yy, vv in zip(yy_, v):
            inside = vv > 0.9
            ax.text(vv - 0.01 if inside else vv + 0.01, yy, f"{vv:.0%}", va="center", ha="right" if inside else "left",
                    fontsize=8, color="white" if inside else INK)
    ax.set_yticks(y)
    ax.set_yticklabels([f"{r.req_id} {str(r.requirement)[:38]}" + ("" if r.critical else " (optional)") for _, r in fz.iterrows()])
    ax.set_xlim(0, 1)
    ax.axvline(0.5, color=GRIDC, lw=1.2)
    ax.set_xlabel("Probability the requirement is met")
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.legend(frameon=False, fontsize=8.5, loc="lower right")
    sub = " · ".join(f"+{r.horizon_years}y: {r.verdict} (joint {r.joint_if_independent:.0%}, weakest {r.weakest_requirement} "
                     f"{r.weakest_link_bound:.0%})" for _, r in ps_df.iterrows() if "verdict" in r and isinstance(r.verdict, str))
    fig.text(0.01, 0.985, a.title or "Project feasibility", fontsize=13, fontweight="bold", color=INK, va="top")
    fig.text(0.01, 0.94, sub, fontsize=8.6, color=MUTED, va="top", wrap=True)
    fig.subplots_adjust(top=0.86)
    fig.savefig(os.path.join(a.outdir, "plots", "feasibility.png"), dpi=150, bbox_inches="tight", pad_inches=0.3,
                facecolor="white")
    plt.close(fig)


# ---------------------------------------------------------------- workbook
SHEETS = [
    ("Project score", "project_score.csv", "Verdict per horizon. Joint = all critical requirements met, if they were independent "
     "(an estimate); weakest link = the best the project can do, since it cannot beat its worst requirement."),
    ("Feasibility", "feasibility.csv", "One row per requirement. p_Ny = probability it is met N years after the origin year, "
     "with ranges widened beyond 5 years for physical technologies (p_model = without widening)."),
    ("Series fit", "series_fit.csv", "One row per history series: trend per year, volatility, grade and whether the trend "
     "has bent (break check)."),
    ("Forecasts", "forecasts.csv", "Forecast ranges by year. q500 = median; q100/q900 = 80% range; q025/q975 = 95% range."),
    ("Calibration", "calibration.csv", "How well the method would have done on this data in the past: share of real "
     "outcomes inside the 80% and 95% ranges. Near 0.80 and 0.95 = honest ranges."),
    ("Events", "events.csv", "Dated future events (laws, patent expiries, factories, launches) and their sources."),
    ("Signposts", "signposts.csv", "Early signs to watch, and the date to re-run the forecast."),
    ("Predictions log", "predictions_log.csv", "Every dated, checkable claim this run makes, to score later."),
    ("Requirements", "requirements.csv", "What the project needs, as written for this run."),
    ("Points", "points.csv", "Every data point with its source."),
]


def cmd_workbook(a):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("workbook needs openpyxl: pip install openpyxl")
    F = "Arial"
    head = PatternFill("solid", fgColor="1F2430")
    zebra = PatternFill("solid", fgColor="F3F4F7")
    side = Side(style="thin", color="9CA0AA")
    bd = Border(left=side, right=side, top=side, bottom=side)
    wb = Workbook()
    ws = wb.active
    ws.title = "Start here"
    lines = [(f"TREND-PREDICTOR: {a.project}", True), (f"Field: {a.field}   Origin year: {a.end}   Run date: {a.date}", False), ("", False),
             ("How to read this in 5 minutes", True),
             ("1. 'Project score' gives the verdict for each horizon and the weakest requirement.", False),
             ("2. 'Feasibility' shows each requirement's chance of being met by each horizon.", False),
             ("3. 'Calibration' tells you whether to trust the ranges: values near 0.80 / 0.95 are honest.", False),
             ("4. 'Signposts' lists what to watch; re-run on the date given.", False), ("", False),
             ("Glossary", True),
             ("Median forecast: half of plausible futures are above, half below.", False),
             ("80% range: 8 in 10 plausible futures fall inside it.", False),
             ("Volatility: how jumpy the history is; jumpier history means wider ranges.", False),
             ("Widening: beyond 5 years, ranges for physical technologies are made wider because long forecasts miss more.", False),
             ("Weakest link: the requirement least likely to be met; the project cannot do better than it.", False),
             ("TRL: technology readiness level, 1 (idea) to 9 (proven in real use).", False)]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(name=F, size=12 if b else 10, bold=b)
    ws.column_dimensions["A"].width = 110
    for name, f, intro in SHEETS:
        df = R(os.path.join(a.dir, f))
        s = wb.create_sheet(name[:31])
        s.cell(1, 1, intro).font = Font(name=F, size=10, italic=True, color="4B4B55")
        if df.empty:
            s.cell(3, 1, "No data for this sheet in this run.").font = Font(name=F, size=10, color="4B4B55")
            continue
        for j, col in enumerate(df.columns, 1):
            c = s.cell(3, j, col)
            c.font, c.fill, c.border = Font(name=F, size=10, bold=True, color="FFFFFF"), head, bd
            c.alignment = Alignment(wrap_text=True, vertical="center")
            width = min(60, max(10, int(df[col].astype(str).str.len().quantile(0.9) if len(df) else 10) + 2, len(col) + 2))
            s.column_dimensions[get_column_letter(j)].width = width
        for i, (_, row) in enumerate(df.iterrows(), 4):
            for j, col in enumerate(df.columns, 1):
                v = row[col]
                c = s.cell(i, j)
                c.border, c.font = bd, Font(name=F, size=10)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if i % 2:
                    c.fill = zebra
                fv = num(v)
                if str(v).startswith("http"):
                    c.value, c.hyperlink = v, v
                    c.font = Font(name=F, size=10, color="1D4ED8", underline="single")
                elif fv == fv and re.fullmatch(r"-?[\d.,eE+-]+", str(v).strip() or "x"):
                    c.value = fv
                else:
                    c.value = v
        s.freeze_panes = "A4"
        s.auto_filter.ref = f"A3:{get_column_letter(len(df.columns))}{len(df) + 3}"
    out = os.path.join(a.dir, f"trend_predictor_{re.sub(r'[^a-z0-9]+', '-', a.project.lower()).strip('-')}.xlsx")
    wb.save(out)
    print(f"wrote {out}")


# -------------------------------------------------------------------- main
def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = p.add_subparsers(dest="cmd", required=True)

    def common(c):
        c.add_argument("--points", default="points.csv")
        c.add_argument("--constraints", default="constraints.csv")
        c.add_argument("--cpi", default="cpi.csv")
        c.add_argument("--base-year", type=int)
        c.add_argument("--end", type=int, required=True, help="origin year: last full calendar year of history")
        c.add_argument("--drift", choices=["endpoint", "theilsen", "robust"], default="endpoint",
                       help="endpoint = Farmer & Lafond (validated); theilsen = robust to an outlying first year")
        c.add_argument("--min-diffs", type=int, default=MIN_DIFFS)
        c.add_argument("--no-autofix", action="store_true", help="disable the automatic repairs (frontier, window, gap)")
        c.add_argument("--outdir", default=".")

    c = sp.add_parser("cpi"); c.add_argument("--out", default="cpi.csv"); c.set_defaults(f=cmd_cpi)
    c = sp.add_parser("forecast"); common(c)
    c.add_argument("--horizons", default="5,10")
    c.add_argument("--scenarios", default="scenarios.csv")
    c.add_argument("--requirements", default="requirements.csv", help="draws targets on the fan charts")
    c.add_argument("--plots", action="store_true")
    c.set_defaults(f=cmd_forecast)
    c = sp.add_parser("hindcast"); common(c)
    c.add_argument("--max-tau", type=int, default=10)
    c.set_defaults(f=cmd_hindcast)
    c = sp.add_parser("feasibility"); common(c)
    c.add_argument("--requirements", default="requirements.csv")
    c.add_argument("--events", default="events.csv")
    c.add_argument("--scenarios", default="scenarios.csv")
    c.add_argument("--horizons", default="5,10")
    c.add_argument("--domain", default="hardware", help="default domain for widening and readiness priors")
    c.add_argument("--widen", type=float, default=None, help=f"long-range widening for physical domains (default {WIDEN['physical']})")
    c.add_argument("--lead-years", type=float, default=2, help="years needed to build the project once constraints allow")
    c.add_argument("--since-offset", type=float, default=0, help="years already spent at the current TRL")
    c.add_argument("--title", default="")
    c.add_argument("--plots", action="store_true")
    c.set_defaults(f=cmd_feasibility)
    c = sp.add_parser("workbook")
    c.add_argument("--dir", default=".")
    c.add_argument("--project", required=True)
    c.add_argument("--field", default="")
    c.add_argument("--end", type=int, required=True)
    c.add_argument("--date", required=True)
    c.set_defaults(f=cmd_workbook)
    a = p.parse_args()
    if hasattr(a, "outdir"):
        os.makedirs(a.outdir, exist_ok=True)
    a.f(a)


if __name__ == "__main__":
    main()
